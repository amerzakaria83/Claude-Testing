#!/usr/bin/env python3
"""
UAE Lottery "Lucky Day" analyser and ticket generator.

Game format (as recorded in the spreadsheet):
  * 6 "Day" numbers drawn from 1..31 (no repeats)
  * 1 "Lucky Month" number drawn from 1..12

What this script does
  1. Loads and validates the results spreadsheet.
  2. Looks for patterns (frequency, hot/cold, overdue numbers, sums,
     odd/even split, repeats, pairs, month bias) and tests each one
     against what pure chance would produce.
  3. Back-tests several prediction strategies walk-forward (each draw is
     "predicted" using only the draws before it) and compares them with
     random picks.
  4. Generates tickets for the next draw.

An honest note: a fair lottery draw has no memory. Every combination of
6 days + 1 month has the same chance, 1 in 8,835,372, whatever happened
before. This script can tell you whether the past data looks fair and
can generate tickets, but no strategy here (or anywhere) can raise the
odds of a fair draw. The back-test section shows this with your data.

Usage
  python lucky_day_predictor.py                       # full report + 5 tickets
  python lucky_day_predictor.py --tickets 10 --strategy hot
  python lucky_day_predictor.py --only-predict --seed 7
  python lucky_day_predictor.py path/to/other_results.xlsx

Requires: openpyxl (pip install openpyxl)
"""

from __future__ import annotations

import argparse
import datetime as dt
import itertools
import math
import random
import statistics
import sys
from collections import Counter
from dataclasses import dataclass
from pathlib import Path

DAY_MAX = 31          # day numbers are 1..31
DAY_PICK = 6          # six day numbers per draw
MONTH_MAX = 12        # lucky month is 1..12
DEFAULT_FILE = Path(__file__).parent / "data" / "UAE_Lottery_Lucky_Day_Results.xlsx"
STRATEGIES = ("random", "hot", "cold", "overdue", "weighted", "balanced")


# --------------------------------------------------------------------------
# Loading
# --------------------------------------------------------------------------

@dataclass(frozen=True)
class Draw:
    draw_no: str
    date: dt.date
    days: tuple[int, ...]   # sorted, 6 numbers
    month: int


def load_draws(path: Path) -> list[Draw]:
    """Read the results sheet and return draws oldest-first."""
    try:
        import openpyxl
    except ImportError:
        sys.exit("openpyxl is required: pip install openpyxl")

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows)]

    def col(name: str) -> int:
        try:
            return header.index(name)
        except ValueError:
            sys.exit(f"Column '{name}' not found in {path.name}; got {header}")

    i_no, i_date, i_month = col("Draw No."), col("Date"), col("Lucky Month")
    i_days = [col(f"Day {k}") for k in range(1, DAY_PICK + 1)]

    draws: dict[str, Draw] = {}
    for line_no, row in enumerate(rows, start=2):
        if row is None or row[i_no] is None:
            continue
        try:
            days = tuple(sorted(int(row[i]) for i in i_days))
            month = int(row[i_month])
        except (TypeError, ValueError):
            print(f"  ! skipping row {line_no}: non-numeric values", file=sys.stderr)
            continue
        if len(set(days)) != DAY_PICK or not all(1 <= d <= DAY_MAX for d in days):
            print(f"  ! skipping row {line_no}: bad day numbers {days}", file=sys.stderr)
            continue
        if not 1 <= month <= MONTH_MAX:
            print(f"  ! skipping row {line_no}: bad month {month}", file=sys.stderr)
            continue
        date = row[i_date]
        date = date.date() if isinstance(date, dt.datetime) else date
        draw_no = str(row[i_no]).strip()
        draws[draw_no] = Draw(draw_no, date, days, month)   # de-duplicates
    wb.close()
    return sorted(draws.values(), key=lambda d: (d.date, d.draw_no))


# --------------------------------------------------------------------------
# Random-draw simulation (used as the "pure chance" yardstick)
# --------------------------------------------------------------------------

def random_draw(rng: random.Random) -> tuple[tuple[int, ...], int]:
    return tuple(sorted(rng.sample(range(1, DAY_MAX + 1), DAY_PICK))), rng.randint(1, MONTH_MAX)


def chi_square(counts: Counter, categories: range, expected: float) -> float:
    return sum((counts.get(c, 0) - expected) ** 2 / expected for c in categories)


def simulate_null(n_draws: int, sims: int, rng: random.Random) -> dict[str, list[float]]:
    """Run `sims` fake histories of `n_draws` fair draws; collect test statistics."""
    out = {k: [] for k in ("chi_days", "chi_month", "max_day", "min_day", "max_month",
                           "top_pair", "mean_repeat")}
    exp_day = n_draws * DAY_PICK / DAY_MAX
    exp_month = n_draws / MONTH_MAX
    for _ in range(sims):
        day_c, mon_c, pair_c = Counter(), Counter(), Counter()
        prev, repeats = None, []
        for _ in range(n_draws):
            days, month = random_draw(rng)
            day_c.update(days)
            mon_c[month] += 1
            pair_c.update(itertools.combinations(days, 2))
            if prev is not None:
                repeats.append(len(set(days) & set(prev)))
            prev = days
        out["chi_days"].append(chi_square(day_c, range(1, DAY_MAX + 1), exp_day))
        out["chi_month"].append(chi_square(mon_c, range(1, MONTH_MAX + 1), exp_month))
        out["max_day"].append(max(day_c.get(d, 0) for d in range(1, DAY_MAX + 1)))
        out["min_day"].append(min(day_c.get(d, 0) for d in range(1, DAY_MAX + 1)))
        out["max_month"].append(max(mon_c.get(m, 0) for m in range(1, MONTH_MAX + 1)))
        out["top_pair"].append(max(pair_c.values()))
        out["mean_repeat"].append(statistics.mean(repeats) if repeats else 0)
    return out


def p_upper(sim: list[float], observed: float) -> float:
    """Share of fair simulations at least as extreme (high) as what we saw."""
    return (sum(s >= observed for s in sim) + 1) / (len(sim) + 1)


def p_lower(sim: list[float], observed: float) -> float:
    return (sum(s <= observed for s in sim) + 1) / (len(sim) + 1)


def verdict(p: float) -> str:
    if p < 0.01:
        return "UNUSUAL - worth checking (could be a data error)"
    if p < 0.05:
        return "slightly unusual, but expected now and then by chance"
    return "normal for random draws"


# --------------------------------------------------------------------------
# Pattern analysis
# --------------------------------------------------------------------------

def bar(n: int, scale: float = 1.0) -> str:
    return "#" * round(n * scale)


def analyse(draws: list[Draw], sims: int, rng: random.Random) -> None:
    n = len(draws)
    print("=" * 72)
    print(f"DATA: {n} draws, {draws[0].date} -> {draws[-1].date}")
    print("=" * 72)
    gaps = [(b.date - a.date).days for a, b in zip(draws, draws[1:])]
    big_gaps = [(a.date, b.date) for a, b in zip(draws, draws[1:]) if (b.date - a.date).days > 15]
    print(f"Typical gap between draws: {statistics.median(gaps):.0f} days; "
          f"{len(big_gaps)} gaps longer than 15 days (missing draws).")
    print("Missing draws do not bias the tests below: fair draws are independent.\n")

    print(f"Simulating {sims:,} fair lottery histories of {n} draws for comparison...\n")
    null = simulate_null(n, sims, rng)

    # ---- 1. Day number frequency --------------------------------------
    day_c = Counter(d for dr in draws for d in dr.days)
    exp_day = n * DAY_PICK / DAY_MAX
    last_seen = {}
    for idx, dr in enumerate(draws):
        for d in dr.days:
            last_seen[d] = idx
    print("-" * 72)
    print(f"1. DAY NUMBER FREQUENCY  (each number expected ~{exp_day:.1f} times)")
    print("-" * 72)
    print(" num  hits  vs exp  draws since last seen")
    for d in range(1, DAY_MAX + 1):
        hits = day_c.get(d, 0)
        since = (n - 1 - last_seen[d]) if d in last_seen else n
        print(f" {d:>3}  {hits:>4}  {hits - exp_day:+6.1f}  {since:>3}  {bar(hits)}")

    chi = chi_square(day_c, range(1, DAY_MAX + 1), exp_day)
    hot = sorted(range(1, DAY_MAX + 1), key=lambda d: -day_c.get(d, 0))[:6]
    cold = sorted(range(1, DAY_MAX + 1), key=lambda d: day_c.get(d, 0))[:6]
    hi, lo = day_c[hot[0]], day_c.get(cold[0], 0)
    print(f"\n Hottest: {hot}   Coldest: {cold}")
    p = p_upper(null["chi_days"], chi)
    print(f" Overall spread (chi-square {chi:.1f}): p = {p:.3f} -> {verdict(p)}")
    p = p_upper(null["max_day"], hi)
    print(f" Most-drawn number appears {hi}x: p = {p:.3f} -> {verdict(p)}")
    p = p_lower(null["min_day"], lo)
    print(f" Least-drawn number appears {lo}x: p = {p:.3f} -> {verdict(p)}")

    # ---- 2. Lucky Month -------------------------------------------------
    mon_c = Counter(dr.month for dr in draws)
    exp_mon = n / MONTH_MAX
    print("\n" + "-" * 72)
    print(f"2. LUCKY MONTH FREQUENCY  (each month expected ~{exp_mon:.1f} times)")
    print("-" * 72)
    for m in range(1, MONTH_MAX + 1):
        print(f" {m:>3}  {mon_c.get(m, 0):>4}  {bar(mon_c.get(m, 0))}")
    top_m = max(range(1, MONTH_MAX + 1), key=lambda m: mon_c.get(m, 0))
    chi = chi_square(mon_c, range(1, MONTH_MAX + 1), exp_mon)
    p = p_upper(null["chi_month"], chi)
    print(f"\n Overall spread (chi-square {chi:.1f}): p = {p:.3f} -> {verdict(p)}")
    p = p_upper(null["max_month"], mon_c[top_m])
    print(f" Month {top_m} drawn {mon_c[top_m]}x (expected {exp_mon:.1f}): "
          f"p = {p:.3f} -> {verdict(p)}")
    print("   (This p-value already allows for the fact that we looked at all 12 months.)")

    # ---- 3. Shape of each draw ----------------------------------------
    print("\n" + "-" * 72)
    print("3. SHAPE OF EACH DRAW")
    print("-" * 72)
    sums = [sum(dr.days) for dr in draws]
    # Exact mean/sd of the sum of 6 numbers drawn without replacement from 1..31
    pop_var = (DAY_MAX ** 2 - 1) / 12
    exp_sum = DAY_PICK * (DAY_MAX + 1) / 2
    sd_sum = math.sqrt(DAY_PICK * pop_var * (DAY_MAX - DAY_PICK) / (DAY_MAX - 1))
    print(f" Sum of the 6 days: mean {statistics.mean(sums):.1f} (random: {exp_sum:.0f}), "
          f"sd {statistics.stdev(sums):.1f} (random: {sd_sum:.1f}), "
          f"range {min(sums)}-{max(sums)}")

    odd_c = Counter(sum(d % 2 for d in dr.days) for dr in draws)
    n_odd = (DAY_MAX + 1) // 2
    total = math.comb(DAY_MAX, DAY_PICK)
    print(" Odd numbers per draw   seen   random would give")
    for k in range(DAY_PICK + 1):
        prob = math.comb(n_odd, k) * math.comb(DAY_MAX - n_odd, DAY_PICK - k) / total
        print(f"   {k} odd / {DAY_PICK - k} even        {odd_c.get(k, 0):>3}    {prob * n:5.1f}")

    consec = sum(any(b - a == 1 for a, b in zip(dr.days, dr.days[1:])) for dr in draws)
    p_no_consec = math.comb(DAY_MAX - DAY_PICK + 1, DAY_PICK) / total
    print(f" Draws containing at least one consecutive pair (e.g. 16,17): "
          f"{consec}/{n} (random: {(1 - p_no_consec) * n:.1f})")

    # ---- 4. Repeats and pairs ----------------------------------------
    print("\n" + "-" * 72)
    print("4. REPEATS AND PAIRS")
    print("-" * 72)
    repeats = [len(set(a.days) & set(b.days)) for a, b in zip(draws, draws[1:])]
    mean_rep = statistics.mean(repeats)
    exp_rep = DAY_PICK * DAY_PICK / DAY_MAX
    p = p_upper(null["mean_repeat"], mean_rep)
    print(f" Numbers repeated from the previous draw: average {mean_rep:.2f} "
          f"(random: {exp_rep:.2f}) -> {verdict(min(p, p_lower(null['mean_repeat'], mean_rep)) * 2)}")
    pair_c = Counter(p for dr in draws for p in itertools.combinations(dr.days, 2))
    exp_pair = n * (DAY_PICK * (DAY_PICK - 1)) / (DAY_MAX * (DAY_MAX - 1))
    top = pair_c.most_common(5)
    print(f" Most common pairs (each pair expected ~{exp_pair:.1f}x): "
          + ", ".join(f"{a}-{b} ({c}x)" for (a, b), c in top))
    p = p_upper(null["top_pair"], top[0][1])
    print(f" Top pair count {top[0][1]}: p = {p:.3f} -> {verdict(p)}")
    print(f" Note: there are {math.comb(DAY_MAX, 2)} possible pairs, so some always "
          "look 'popular' by luck.")


# --------------------------------------------------------------------------
# Strategies: each returns a weight per day number and per month
# --------------------------------------------------------------------------

def weights_for(strategy: str, history: list[Draw]) -> tuple[dict[int, float], dict[int, float]]:
    days = range(1, DAY_MAX + 1)
    months = range(1, MONTH_MAX + 1)
    day_c = Counter(d for dr in history for d in dr.days)
    mon_c = Counter(dr.month for dr in history)
    n = len(history)

    if strategy == "hot":           # favour numbers drawn most often
        return ({d: (day_c[d] + 0.5) ** 2 for d in days},
                {m: (mon_c[m] + 0.5) ** 2 for m in months})
    if strategy == "cold":          # favour numbers drawn least often
        top_d, top_m = max(day_c.values(), default=0), max(mon_c.values(), default=0)
        return ({d: (top_d - day_c[d] + 0.5) ** 2 for d in days},
                {m: (top_m - mon_c[m] + 0.5) ** 2 for m in months})
    if strategy == "overdue":       # favour numbers not seen for a long time
        last_d, last_m = {}, {}
        for i, dr in enumerate(history):
            for d in dr.days:
                last_d[d] = i
            last_m[dr.month] = i
        return ({d: (n - last_d.get(d, -1)) ** 1.5 for d in days},
                {m: (n - last_m.get(m, -1)) ** 1.5 for m in months})
    if strategy == "weighted":      # mild lean towards recent frequency
        dw = {d: 1.0 for d in days}
        mw = {m: 1.0 for m in months}
        for age, dr in enumerate(reversed(history)):
            decay = 0.97 ** age
            for d in dr.days:
                dw[d] += decay
            mw[dr.month] += decay
        return dw, mw
    # "random" and "balanced": every number equal
    return {d: 1.0 for d in days}, {m: 1.0 for m in months}


def weighted_sample(weights: dict[int, float], k: int, rng: random.Random) -> list[int]:
    pool = dict(weights)
    picked = []
    for _ in range(k):
        total = sum(pool.values())
        r = rng.uniform(0, total)
        acc = 0.0
        for num, w in pool.items():
            acc += w
            if acc >= r:
                break
        picked.append(num)
        del pool[num]
    return sorted(picked)


def looks_typical(days: list[int]) -> bool:
    """Shape filter used by 'balanced': middle-of-the-road sum, 2-4 odd numbers,
    not a run of 3+ consecutive numbers. This makes tickets look like typical
    draws. It does NOT change the odds of winning - see README."""
    odd = sum(d % 2 for d in days)
    runs = any(days[i + 2] - days[i] == 2 for i in range(len(days) - 2))
    return 72 <= sum(days) <= 120 and 2 <= odd <= 4 and not runs


def make_ticket(strategy: str, history: list[Draw], rng: random.Random) -> tuple[list[int], int]:
    dw, mw = weights_for(strategy, history)
    for _ in range(1000):
        days = weighted_sample(dw, DAY_PICK, rng)
        if strategy != "balanced" or looks_typical(days):
            break
    month = weighted_sample(mw, 1, rng)[0]
    return days, month


# --------------------------------------------------------------------------
# Back-test
# --------------------------------------------------------------------------

def backtest(draws: list[Draw], rng: random.Random, start: int = 10, reps: int = 200) -> None:
    print("\n" + "-" * 72)
    print("5. BACK-TEST: would any strategy have predicted the past draws?")
    print("-" * 72)
    targets = draws[start:]
    print(f" For each of the last {len(targets)} draws, generate {reps} tickets per strategy")
    print(" using ONLY the draws before it, then count matches.\n")
    exp_match = DAY_PICK * DAY_PICK / DAY_MAX
    print(f" {'strategy':<10} {'avg days matched':>17} {'month hit rate':>15} {'3+ days hit':>12}")
    print(f" {'(pure luck)':<10} {exp_match:>17.3f} {1 / MONTH_MAX:>15.1%} {p_at_least(3):>12.2%}")
    for strat in STRATEGIES:
        matches, month_hits, three_plus, total = 0, 0, 0, 0
        for i, target in enumerate(targets, start=start):
            history = draws[:i]
            for _ in range(reps):
                days, month = make_ticket(strat, history, rng)
                m = len(set(days) & set(target.days))
                matches += m
                three_plus += m >= 3
                month_hits += month == target.month
                total += 1
        print(f" {strat:<10} {matches / total:>17.3f} {month_hits / total:>15.1%} "
              f"{three_plus / total:>12.2%}")
    print("\n Small differences between rows are noise: with this few draws, "
          "+/-0.2 matches\n or +/-4% on the month happen by chance alone.")


def p_at_least(k: int) -> float:
    total = math.comb(DAY_MAX, DAY_PICK)
    return sum(math.comb(DAY_PICK, j) * math.comb(DAY_MAX - DAY_PICK, DAY_PICK - j)
               for j in range(k, DAY_PICK + 1)) / total


# --------------------------------------------------------------------------
# Main
# --------------------------------------------------------------------------

def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("file", nargs="?", type=Path, default=DEFAULT_FILE, help="results .xlsx")
    ap.add_argument("--tickets", type=int, default=5, help="how many tickets to generate")
    ap.add_argument("--strategy", choices=STRATEGIES, default="balanced")
    ap.add_argument("--seed", type=int, help="random seed for repeatable output")
    ap.add_argument("--sims", type=int, default=5000, help="fair-lottery simulations for p-values")
    ap.add_argument("--only-predict", action="store_true", help="skip analysis and back-test")
    args = ap.parse_args()

    rng = random.Random(args.seed)
    draws = load_draws(args.file)
    if len(draws) < 12:
        sys.exit(f"Need at least 12 valid draws, found {len(draws)}.")

    if not args.only_predict:
        analyse(draws, args.sims, rng)
        backtest(draws, rng)

    odds = math.comb(DAY_MAX, DAY_PICK) * MONTH_MAX
    print("\n" + "=" * 72)
    print(f"TICKETS FOR THE NEXT DRAW  (strategy: {args.strategy})")
    print("=" * 72)
    seen = set()
    while len(seen) < args.tickets:
        days, month = make_ticket(args.strategy, draws, rng)
        key = (tuple(days), month)
        if key in seen:
            continue
        seen.add(key)
        print(f" Ticket {len(seen)}:  Days {' '.join(f'{d:>2}' for d in days)}   "
              f"Lucky Month {month:>2}")
    print(f"\n Chance that any one ticket hits all 6 days + month: 1 in {odds:,}.")
    print(" That is the same for every ticket and every strategy. Play for fun, "
          "with money you can afford to lose.")


if __name__ == "__main__":
    main()
